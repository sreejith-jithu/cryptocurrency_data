import requests
import pandas as pd 
import time
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, Alignment, PatternFill


EXCEL_FILE = "crypto_data.xlsx"  # Excel file name
UPDATE_INTERVAL = 300  # 5 minutes (300 seconds)

# Function to fetch live cryptocurrency data
def fetch_crypto_data():
    url = "https://api.coingecko.com/api/v3/coins/markets"
    params = {
        "vs_currency": "usd",
        "order": "market_cap_desc",
        "per_page": 50,
        "page": 1,
        "sparkline": False
    }
    response = requests.get(url, params=params)

    if response.status_code != 200:
        print(f"Error: Unable to fetch data. Status Code: {response.status_code}")
        return pd.DataFrame()  # Return empty DataFrame if request fails

    data = response.json()

    # Extract relevant fields
    crypto_list = []
    for coin in data:
        crypto_list.append({
            "Name": coin["name"],
            "Symbol": coin["symbol"].upper(),
            "Price (USD)": coin["current_price"],
            "Market Cap (USD)": coin["market_cap"],
            "24h Volume (USD)": coin["total_volume"],
            "24h Change (%)": coin["price_change_percentage_24h"]
        })
    
    return pd.DataFrame(crypto_list)

# Function to analyze data
def analyze_data(df):
    if df.empty:
        return {}

    analysis = {}

    # Top 5 cryptocurrencies by market cap
    top_5 = df.nlargest(5, "Market Cap (USD)")[["Name", "Market Cap (USD)"]]

    # Average price of the top 50 cryptocurrencies
    avg_price = df["Price (USD)"].mean()

    # Highest and lowest 24h percentage change
    highest_change = df.loc[df["24h Change (%)"].idxmax()]
    lowest_change = df.loc[df["24h Change (%)"].idxmin()]

    analysis["Top 5 by Market Cap"] = top_5
    analysis["Average Price"] = avg_price
    analysis["Highest 24h Change"] = {
        "Name": highest_change["Name"],
        "Change (%)": highest_change["24h Change (%)"]
    }
    analysis["Lowest 24h Change"] = {
        "Name": lowest_change["Name"],
        "Change (%)": lowest_change["24h Change (%)"]
    }

    return analysis

# Function to update Excel file with better formatting
def update_excel(df, analysis):
    if df.empty:
        print("No data to update.")
        return

    try:
        wb = load_workbook(EXCEL_FILE)
    except FileNotFoundError:
        wb = Workbook()

    # Update main data sheet
    if "Crypto Data" not in wb.sheetnames:
        ws = wb.create_sheet("Crypto Data")
    else:
        ws = wb["Crypto Data"]
        ws.delete_rows(2, ws.max_row)  # Clear old data

    # Add header if it's empty
    if ws.max_row == 1:
        ws.append(["Name", "Symbol", "Price (USD)", "Market Cap (USD)", "24h Volume (USD)", "24h Change (%)"])

    # Add new data
    for _, row in df.iterrows():
        ws.append(row.tolist())

    # Update analysis sheet
    if "Analysis" not in wb.sheetnames:
        ws_analysis = wb.create_sheet("Analysis")
    else:
        ws_analysis = wb["Analysis"]
        ws_analysis.delete_rows(2, ws_analysis.max_row)  # Clear old data

    ws_analysis.append(["Metric", "Value"])
    ws_analysis.append(["Average Price of Top 50", f"${analysis['Average Price']:.2f}"])

    ws_analysis.append(["Top 5 Cryptos by Market Cap"])
    for _, row in analysis["Top 5 by Market Cap"].iterrows():
        ws_analysis.append([row["Name"], f"${row['Market Cap (USD)']:,}"])

    ws_analysis.append(["Highest 24h Change", analysis["Highest 24h Change"]["Name"], f"{analysis['Highest 24h Change']['Change (%)']:.2f}%"])
    ws_analysis.append(["Lowest 24h Change", analysis["Lowest 24h Change"]["Name"], f"{analysis['Lowest 24h Change']['Change (%)']:.2f}%"])

    # Apply formatting
    format_excel(wb)

    # Save the Excel file
    wb.save(EXCEL_FILE)
    
    print(f"✅ Excel file '{EXCEL_FILE}' updated successfully.")

# Function to format the Excel file
def format_excel(wb):
    ws_data = wb["Crypto Data"]
    ws_analysis = wb["Analysis"]

    # Formatting: Bold Headers
    for cell in ws_data[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center")
        cell.fill = PatternFill(start_color="4F81BD", end_color="4F81BD", fill_type="solid")  # Blue header

    for cell in ws_analysis[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")

    # Auto-adjust column width
    for ws in [ws_data, ws_analysis]:
        for col in ws.columns:
            max_length = max((len(str(cell.value)) for cell in col if cell.value), default=10)
            ws.column_dimensions[col[0].column_letter].width = max_length + 2

    # Format numbers in Crypto Data sheet
    for row in ws_data.iter_rows(min_row=2, min_col=3, max_col=6):
        row[0].number_format = "$#,##0.00"   # Price
        row[1].number_format = "$#,##0"      # Market Cap
        row[2].number_format = "$#,##0"      # 24h Volume
        row[3].number_format = "0.00%"       # 24h Change

    # Highlight key metrics in Analysis sheet
    highlight_fill = PatternFill(start_color="FFFF99", end_color="FFFF99", fill_type="solid")  # Yellow fill
    for row in ws_analysis.iter_rows(min_row=2, max_col=2):
        for cell in row:
            cell.fill = highlight_fill
            cell.alignment = Alignment(horizontal="center")

# Main loop to fetch and update data every 5 minutes
if __name__ == "__main__":
    while True:
        print("📡 Fetching live cryptocurrency data...")
        df = fetch_crypto_data()

        if df.empty:
            print("⚠ No data received. Retrying in 5 minutes...")
            time.sleep(UPDATE_INTERVAL)
            continue

        print(df.head())  
        analysis = analyze_data(df)
        update_excel(df, analysis)
        print("⏳ Waiting for the next update...\n")
        time.sleep(UPDATE_INTERVAL)  